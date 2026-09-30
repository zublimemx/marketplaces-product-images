#!/usr/bin/env python3
"""Versiona los ajustes de precio exportados desde el visor («Exportar ajustes de precio») en data/ajustes_precios.json.

Uso:
  python scripts/ajustes_precios.py ajustes_precios_2026-09-29_2140.xlsx     # aplica el archivo
  python scripts/ajustes_precios.py --lista                                   # muestra los ajustes versionados
  python scripts/ajustes_precios.py --quitar 7501... [7502...]                # borra los ajustes de esos productos

Cada renglón del archivo reemplaza los ajustes de ese producto: las celdas con valor se guardan y las vacías dejan
de estar ajustadas (se vuelve al valor del sistema, al envío estimado o al dato de la API). La comisión puede venir
como 0.14 o como 14 (%). Después: regenerar el layout y el visor (docs/PROCEDIMIENTO_SESION.md §4).
"""
import argparse
import datetime
import json
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import precios  # noqa: E402

# Encabezado del Excel del visor → campo
COLUMNAS = {
    "Precio de venta (IVA incluido)": "precio_venta",
    "Costo de empaque y logística": "costo_empaque",
    "Comisión Meli (%, IVA incluido)": "comision",
    "Costo de envío si el precio queda en $299 o más": "costo_envio",
    "Promedio otros vendedores (referencia)": "precio_promedio_otros",
    "Precio mejor vendedor": "precio_mejor_vendedor",
    "Descuento contra mejor vendedor": "descuento_mejor_vendedor",
}


def guardar(aj):
    with open(precios.AJUSTES, "w", encoding="utf-8") as fh:
        json.dump(dict(sorted(aj.items())), fh, ensure_ascii=False, indent=2)
        fh.write("\n")


def aplicar(archivo):
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    ws = wb["Ajustes"] if "Ajustes" in wb.sheetnames else wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    hdr = [str(h or "").strip() for h in next(filas)]
    idx = {h: i for i, h in enumerate(hdr)}
    if "Código" not in idx:
        raise SystemExit("No encuentro la columna «Código»")
    aj = precios.cargar_ajustes()
    hoy = datetime.date.today().isoformat()
    cambiados, quitados = 0, 0
    for r in filas:
        g = r[idx["Código"]]
        if g in (None, ""):
            continue
        g = str(g).strip().removesuffix(".0")
        if not os.path.exists(os.path.join(precios.ROOT, "products", g, "product.json")):
            print(f"Aviso: {g} no existe en products/; se omite")
            continue
        nuevo = {}
        for h, campo in COLUMNAS.items():
            if h in idx and r[idx[h]] not in (None, ""):
                v = float(r[idx[h]])
                if campo == "comision" and v >= 1:
                    v = v / 100
                nuevo[campo] = round(v, 4) if campo == "comision" else round(v, 2)
        nota = r[idx["Nota"]] if "Nota" in idx else None
        if nuevo:
            if nota:
                nuevo["nota"] = str(nota).strip()
            nuevo["fecha"] = hoy
            aj[g] = nuevo
            cambiados += 1
        elif g in aj:
            del aj[g]
            quitados += 1
    guardar(aj)
    print(f"{cambiados} productos con ajustes guardados, {quitados} sin ajustes -> {os.path.relpath(precios.AJUSTES, precios.ROOT)} ({len(aj)} en total)")
    print("Siguiente: regenera el layout y el visor (docs/PROCEDIMIENTO_SESION.md §4).")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("archivo", nargs="?")
    ap.add_argument("--lista", action="store_true")
    ap.add_argument("--quitar", nargs="+")
    a = ap.parse_args()
    if a.lista:
        for g, v in precios.cargar_ajustes().items():
            print(g, json.dumps(v, ensure_ascii=False))
        return
    if a.quitar:
        aj = precios.cargar_ajustes()
        n = sum(1 for g in a.quitar if aj.pop(g, None) is not None)
        guardar(aj)
        print(f"{n} productos sin ajustes")
        return
    if not a.archivo:
        ap.error("indica el archivo exportado del visor, --lista o --quitar")
    aplicar(a.archivo)


if __name__ == "__main__":
    main()
