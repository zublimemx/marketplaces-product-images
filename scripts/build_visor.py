#!/usr/bin/env python3
"""Genera los datos del visor de productos (visor/data/productos.js).

Uso:
  python scripts/build_visor.py [--precios insumos/precios_existencias.csv] [--competencia insumos/competencia_meli.csv]
                                [--base-imagenes ../products] [--salida visor/data/productos.js]

Une product.json, precios y existencias, competencia de Mercado Libre y los precios calculados con las mismas
fórmulas del layout (scripts/precios.py), y califica cada producto con los indicadores de config/indicadores.json.
Si falta el CSV de precios, usa la hoja Precios del layout versionado (layouts/mercadolibre/layout_mercadolibre.xlsx).
El archivo es JavaScript (window.CATALOGO = …) para que el visor abra también con doble clic, sin servidor.
"""
import argparse
import csv
import datetime
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import precios  # noqa: E402
import build_mercadolibre as bml  # noqa: E402
from revisar_fotos import fondo_no_blanco  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, "trabajo", "cache_fondo_fotos.json")
SECCIONES = []


def num(v):
    try:
        return float(v) if v not in ("", None) else None
    except (TypeError, ValueError):
        return None


def leer_orden():
    f = os.path.join(ROOT, "data", "prioridad.csv")
    if not os.path.exists(f):
        return {}
    with open(f, encoding="utf-8") as fh:
        return {r["gtin"]: int(r["orden"]) for r in csv.DictReader(fh)}


def leer_precios(path):
    datos = {}
    orden = leer_orden()
    if path and os.path.exists(path):
        with open(path, encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                datos[r["gtin"]] = {"precio": float(r["precio"]), "stock": int(float(r["stock"])), "orden": orden.get(r["gtin"], len(datos) + 1)}
        return datos
    xlsx = os.path.join(ROOT, "layouts", "mercadolibre", "layout_mercadolibre.xlsx")
    if not os.path.exists(xlsx):
        sys.exit("No hay insumos/precios_existencias.csv ni layout versionado para tomar precios y existencias")
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    pre = list(wb["Precios"].iter_rows(min_row=2, values_only=True))
    lay = list(wb["Layout Mercado Libre"].iter_rows(min_row=2, values_only=True))
    for i, (rp, rl) in enumerate(zip(pre, lay), 1):
        datos[str(rp[0])] = {"precio": float(rp[3]), "stock": int(rl[6] or 0), "orden": orden.get(str(rp[0]), i)}
    if "Descartados" in wb.sheetnames:
        for rd in wb["Descartados"].iter_rows(min_row=2, values_only=True):
            if rd[0]:
                datos[str(rd[0])] = {"precio": float(rd[5]), "stock": int(rd[6] or 0), "orden": orden.get(str(rd[0]), len(datos) + 1)}
    return dict(sorted(datos.items(), key=lambda kv: kv[1]["orden"]))


def leer_competencia(path):
    comp = {}
    if path and os.path.exists(path):
        with open(path, encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                comp[r["gtin"]] = (num(r.get("precio_promedio_otros")), num(r.get("precio_mejor_vendedor")), r.get("metodo_mejor_vendedor", ""))
    return comp


def fondo_gris(path, cache):
    st = os.stat(path)
    clave = f"{os.path.relpath(path, ROOT)}|{st.st_size}|{int(st.st_mtime)}"
    if clave not in cache:
        cache[clave] = fondo_no_blanco(path)
    return cache[clave]


def ind_descripcion(p, reglas):
    d = p.get("descripcion") or ""
    inv = p["investigacion"]
    secs = sum(1 for s in reglas["secciones_reconocidas"] if re.search(rf"(^|\n){re.escape(s)}:", d))
    n = len(d)
    base = f"{n:,} caracteres, {secs} secciones"
    if inv["estado"] != reglas["regular"]["estado"]:
        return "mala", ("Sin verificar: datos deducidos del nombre" if inv["estado"] == "sin_verificar" else "Pendiente de investigar")
    if n < reglas["regular"]["min_caracteres"]:
        return "mala", f"Muy corta: {base}"
    b = reglas["buena"]
    if inv.get("confianza") in b["confianza"] and n >= b["min_caracteres"] and secs >= b["min_secciones"]:
        return "buena", f"Verificada con confianza {inv.get('confianza')}; {base}"
    faltas = []
    if inv.get("confianza") not in b["confianza"]:
        faltas.append(f"confianza {inv.get('confianza') or 'sin dato'}")
    if n < b["min_caracteres"]:
        faltas.append(f"menos de {b['min_caracteres']} caracteres")
    if secs < b["min_secciones"]:
        faltas.append(f"menos de {b['min_secciones']} secciones")
    return "regular", f"{base}; " + ", ".join(faltas)


def ind_fotos(imgs, reglas):
    if not imgs:
        return "mala", "Sin fotos"
    utiles = [im["u"] for im in imgs]
    if max(utiles) < reglas["regular"]["min_lado_util"]:
        return "mala", f"{len(imgs)} foto(s), producto de {max(utiles)} px como máximo (mínimo {reglas['regular']['min_lado_util']})"
    b = reglas["buena"]
    principal = imgs[0]
    faltas = []
    if len(imgs) < b["min_fotos"]:
        faltas.append(f"solo {len(imgs)} foto")
    if principal["u"] < b["min_lado_util_principal"]:
        faltas.append(f"principal de {principal['u']} px")
    if b.get("principal_sin_fondo_gris") and principal.get("gris"):
        faltas.append("principal con posible fondo gris")
    if not faltas:
        return "buena", f"{len(imgs)} fotos; principal de {principal['u']} px"
    return "regular", f"{len(imgs)} foto(s); " + ", ".join(faltas)


def ind_precios(pr, reglas):
    falta = [txt for k, txt in reglas["requeridos"].items() if pr.get(k) is None]
    if not falta:
        return "completos", "Se conocen los tres precios"
    return "incompletos", "Falta: " + "; ".join(falta)


FICHA_CLAVE = {"marca": "marca", "fabricante": "fabricante", "presentacion": "presentación",
               "contenido_neto": "contenido neto", "unidad_contenido": "unidad de contenido"}


def _frases(texto, patron=None, maximo=2, largo=280):
    frases = [f.strip() for f in re.split(r"(?<=\.)\s+(?=[A-ZÁÉÍÓÚÑ0-9¿(«\"])", texto or "") if f.strip()]
    if patron is not None:
        frases = [f for f in frases if patron.search(f)]
    out = " ".join(frases[:maximo])
    return out if len(out) <= largo else out[: largo - 1].rstrip() + "…"


def _dinero(v):
    return f"${v:,.2f}"


def pendientes(p, imgs, pr, stock, cfg_pend, cfg_precios):
    """Lista de pendientes, errores y mejoras del producto: [{"c": código, "d": detalle}]. Catálogo en config/pendientes.json."""
    u = cfg_pend["umbrales"]
    out = []
    add = lambda c, d: out.append({"c": c, "d": d})
    inv = p["investigacion"]
    ml = p["marketplaces"]["mercadolibre"]
    notas = inv.get("notas") or ""
    titulo = p.get("titulo") or ""
    desc = p.get("descripcion") or ""
    secs = sum(1 for s in SECCIONES if re.search(rf"(^|\n){re.escape(s)}:", desc))

    # errores
    if not titulo:
        add("sin_titulo", f"Se usaría el nombre del sistema: {p['nombre_sistema']}")
    elif len(titulo) > u["titulo_max_caracteres"]:
        add("titulo_largo", f"{len(titulo)} caracteres: «{titulo}»")
    if not ml.get("categoria_id"):
        add("sin_categoria", "No tiene categoría hoja asignada")
    if len(desc) < u["descripcion_min_caracteres"]:
        add("descripcion_muy_corta", f"{len(desc):,} caracteres (mínimo {u['descripcion_min_caracteres']})")
    chicas = bool(imgs) and max(im["u"] for im in imgs) < u["foto_min_lado_util"]
    if not imgs:
        ni = (inv.get("notas_imagenes") or "").strip()
        add("sin_fotos", "No se encontraron fotos del producto" + (f". {_frases(ni, maximo=1)}" if len(ni) > 15 and not ni.lower().startswith("sin fotos") else ""))
    elif chicas:
        add("fotos_chicas", f"{len(imgs)} foto(s); el producto ocupa {max(im['u'] for im in imgs)} px como máximo (mínimo {u['foto_min_lado_util']})")
    if not pr["precio_venta"] or pr["precio_venta"] <= 0:
        add("sin_precio_venta", "El catálogo del sistema no trae precio")

    # pendientes
    if inv["estado"] != "verificado":
        add("sin_verificar", f"Confianza {inv.get('confianza') or 'baja'}; datos deducidos del nombre del sistema. {_frases(notas, maximo=1)}".strip())
    pat = re.compile(u["patron_dato_por_confirmar"], re.I)
    dudas = _frases(notas + " " + (inv.get("notas_imagenes") or ""), pat)
    if dudas:
        add("dato_por_confirmar", dudas)
    rx = ml.get("categoria_rx_sugerida") or ""
    if p.get("receta_mx") == "Sí":
        add("receta", "Según su principio activo requiere receta en México; se publica como venta libre por indicación del dueño, pero Mercado Libre podría rechazarlo"
            + (f". Categoría con receta sugerida: {rx}" if rx else ""))
    elif p.get("receta_mx") == "Revisar":
        add("receta", "Las fuentes no coinciden sobre si requiere receta; confirmar la condición de venta" + (f". Categoría con receta sugerida: {rx}" if rx else ""))
    if stock <= 0:
        add("sin_existencia", "Se publicaría con 0 piezas")
    if pr["precio_mejor_vendedor"] is None:
        add("sin_mejor_vendedor", f"Mientras no se tenga, se publica al Precio Meli calculado ({_dinero(pr['precio_meli_calculado'])})")
    if pr["precio_meli_final"] >= cfg_precios["umbral_envio_gratis_obligatorio"] and not cfg_precios["costo_envio_vendedor_estimado"]:
        add("envio_sin_costo", f"Precio Meli final de {_dinero(pr['precio_meli_final'])}: Mercado Libre obliga el envío gratis y lo cobra al vendedor; hoy el costo está en $0, así que el margen ({_dinero(pr['margen'])}) está sobreestimado")

    # mejoras
    if inv["estado"] == "verificado" and inv.get("confianza") != "alta":
        add("confianza_media", f"Confianza {inv.get('confianza') or 'sin dato'}. {_frases(notas, maximo=1)}".strip())
    if u["descripcion_min_caracteres"] <= len(desc) < u["descripcion_buena_caracteres"]:
        add("descripcion_corta", f"{len(desc):,} caracteres (recomendado {u['descripcion_buena_caracteres']} o más)")
    if len(desc) >= u["descripcion_min_caracteres"] and secs < u["descripcion_min_secciones"]:
        add("pocas_secciones", f"{secs} secciones reconocidas (recomendado {u['descripcion_min_secciones']} o más)")
    ficha = p.get("ficha") or {}
    faltan = [txt for k, txt in FICHA_CLAVE.items() if ficha.get(k) in ("", None)]
    if faltan:
        add("ficha_incompleta", "Falta: " + ", ".join(faltan))
    if imgs and not chicas:
        if len(imgs) < u["fotos_min"]:
            add("una_foto", f"Solo {len(imgs)} foto; se recomiendan {u['fotos_min']} o más (frente, reverso, ficha)")
        if imgs[0]["u"] < u["foto_principal_min_lado_util"]:
            add("foto_principal_chica", f"El producto ocupa {imgs[0]['u']} px en la foto principal (recomendado {u['foto_principal_min_lado_util']})")
    if imgs and imgs[0].get("gris"):
        add("fondo_gris", "La foto principal podría no tener fondo blanco")
    mv = pr["precio_mejor_vendedor"]
    if mv is not None and pr["precio_meli_calculado"] > mv - cfg_precios["descuento_vs_mejor_vendedor"]:
        add("no_competitivo", f"Calculado {_dinero(pr['precio_meli_calculado'])} contra mejor vendedor {_dinero(mv)}: se publica al calculado, {pr['precio_meli_calculado'] / mv - 1:.0%} arriba")
    if pr["precio_venta"] and pr["precio_meli_final"] / pr["precio_venta"] >= u["veces_precio_tienda"]:
        add("precio_inflado", f"Precio Meli final {_dinero(pr['precio_meli_final'])} = {pr['precio_meli_final'] / pr['precio_venta']:.1f} veces el precio de tienda ({_dinero(pr['precio_venta'])}) por el costo fijo de Mercado Libre; considerar kit o paquete")
    return out


def datos_layout_meli(cfg_precios):
    """Lo que el visor necesita para exportar el layout de Mercado Libre igual que scripts/build_mercadolibre.py."""
    pub = json.load(open(os.path.join(ROOT, "config", "mercadolibre.json"), encoding="utf-8"))["publicacion"]
    return {"encabezados": bml.LAYOUT_HDR, "encabezados_aux": bml.AUX_HDR, "ficha": [k for k, _ in bml.FICHA_COLS],
            "n_imagenes": bml.N_IMG, "url_imagenes": bml.IMG_BASE, "publicacion": pub,
            "umbral_envio_gratis": cfg_precios["umbral_envio_gratis_obligatorio"], "texto_envio_gratis": bml.ENVIO_GRATIS,
            "estados": bml.ESTADOS}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--precios", default=os.path.join(ROOT, "insumos", "precios_existencias.csv"))
    ap.add_argument("--competencia", default=os.path.join(ROOT, "insumos", "competencia_meli.csv"))
    ap.add_argument("--base-imagenes", default="../products", help="ruta o URL base de las fotos, vista desde visor/index.html")
    ap.add_argument("--salida", default=os.path.join(ROOT, "visor", "data", "productos.js"))
    a = ap.parse_args()

    cfg = precios.cargar_config()
    reglas = json.load(open(os.path.join(ROOT, "config", "indicadores.json"), encoding="utf-8"))
    cfg_pend = json.load(open(os.path.join(ROOT, "config", "pendientes.json"), encoding="utf-8"))
    SECCIONES[:] = reglas["descripcion"]["secciones_reconocidas"]
    pv = leer_precios(a.precios)
    comp = leer_competencia(a.competencia)
    cache = json.load(open(CACHE)) if os.path.exists(CACHE) else {}

    productos = []
    for g, base in pv.items():
        p = json.load(open(os.path.join(ROOT, "products", g, "product.json"), encoding="utf-8"))
        ml = p["marketplaces"]["mercadolibre"]
        prom, mejor, metodo = comp.get(g, (None, None, ""))
        pr = precios.calcular(base["precio"], ml["categoria_ruta"], cfg, prom, mejor)
        imgs = []
        for im in p.get("imagenes", []):
            f = os.path.join(ROOT, "products", g, "images", im["archivo"])
            if not os.path.exists(f):
                continue
            imgs.append({"src": f"{a.base_imagenes.rstrip('/')}/{g}/images/{im['archivo']}", "a": im["archivo"], "u": im.get("lado_util", 0),
                         "o": im.get("origen", ""), "gris": fondo_gris(f, cache), "fuente": im.get("pagina") or im.get("fuente_url", "")})
        d_ind, d_mot = ind_descripcion(p, reglas["descripcion"])
        f_ind, f_mot = ind_fotos(imgs, reglas["fotos"])
        p_ind, p_mot = ind_precios(pr, reglas["precios"])
        ruta = ml["categoria_ruta"] or ""
        descartado = ml.get("descartado") or None
        pend = [] if descartado else pendientes(p, imgs, pr, base["stock"], cfg_pend, cfg)
        productos.append({
            "orden": base["orden"], "gtin": g, "titulo": p["titulo"] or p["nombre_sistema"], "nombre_sistema": p["nombre_sistema"],
            "linea": p["linea"], "categoria_id": ml["categoria_id"], "categoria_ruta": ruta,
            "categoria": ruta.split(" > ")[-1] if ruta else "Sin categoría",
            "catalogo_id": ml.get("catalogo_id", ""), "categoria_rx_sugerida": ml.get("categoria_rx_sugerida", ""),
            "stock": base["stock"], "precios": {k: pr[k] for k in ("precio_venta", "precio_marketplaces", "comision", "precio_meli_calculado",
                                                                  "precio_promedio_otros", "precio_mejor_vendedor", "precio_meli_final",
                                                                  "diferencia_mejor_vendedor", "costo_fijo", "envio_vendedor", "ingreso_neto", "margen")},
            "metodo_mejor_vendedor": metodo,
            "imagenes": imgs, "descripcion": p["descripcion"], "ficha": p.get("ficha", {}), "receta_mx": p.get("receta_mx", ""),
            "url_oficial": p.get("url_oficial", ""), "fuentes": p.get("fuentes", []),
            "investigacion": {k: p["investigacion"].get(k, "") for k in ("estado", "confianza", "encontrado_por", "notas", "notas_imagenes", "sesion")},
            "ind": {"descripcion": d_ind, "descripcion_motivo": d_mot, "fotos": f_ind, "fotos_motivo": f_mot, "precios": p_ind, "precios_motivo": p_mot},
            "pend": pend, "descartado": descartado, "sin_titulo": not p["titulo"],
        })
    os.makedirs(os.path.dirname(CACHE), exist_ok=True)
    json.dump(cache, open(CACHE, "w"))
    datos = {"generado": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), "total": len(productos),
             "reglas": reglas, "pendientes": {k: cfg_pend[k] for k in ("acciones", "tipos", "pendientes")},
             "meli": datos_layout_meli(cfg), "productos": productos}
    os.makedirs(os.path.dirname(a.salida), exist_ok=True)
    with open(a.salida, "w", encoding="utf-8") as fh:
        fh.write("// Generado por scripts/build_visor.py. No editar a mano.\nwindow.CATALOGO = ")
        json.dump(datos, fh, ensure_ascii=False, separators=(",", ":"))
        fh.write(";\n")
    cuenta = {}
    for pr_ in productos:
        for k in ("descripcion", "fotos", "precios"):
            cuenta.setdefault(k, {}).setdefault(pr_["ind"][k], 0)
            cuenta[k][pr_["ind"][k]] += 1
    print(f"{len(productos)} productos -> {os.path.relpath(a.salida, ROOT)} ({os.path.getsize(a.salida) / 1e6:.1f} MB)")
    print(json.dumps(cuenta, ensure_ascii=False))
    cp = {}
    for pr_ in productos:
        for x in pr_["pend"]:
            cp[x["c"]] = cp.get(x["c"], 0) + 1
    print("pendientes:", json.dumps(dict(sorted(cp.items(), key=lambda kv: -kv[1])), ensure_ascii=False))
    print("productos con pendientes:", sum(1 for x in productos if x["pend"]), "· descartados:", sum(1 for x in productos if x["descartado"]))


if __name__ == "__main__":
    main()
