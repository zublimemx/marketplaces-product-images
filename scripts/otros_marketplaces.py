#!/usr/bin/env python3
"""Precios de venta de nuestros productos en otros marketplaces y tiendas en línea (Farmacias Benavides, del Ahorro, YZA,
San Pablo, Walmart, Chedraui…), para compararlos con el nuestro. Fuente versionada: data/precios_otros_marketplaces.csv
(un renglón por producto y marketplace). Los lee scripts/build_mercadolibre.py (hoja Precios, columnas Z y AA),
scripts/build_visor.py (pantalla «Comparar precios») y scripts/build_db.py (tabla precios_otros_marketplaces).

Uso:
  python scripts/otros_marketplaces.py integrar --sesion 9       # resultados de subagentes (docs/agentes/precios_otros.md)
  python scripts/otros_marketplaces.py excel <archivo.xlsx>      # hoja «Captura» del Excel que exporta el visor
  python scripts/otros_marketplaces.py lista [--gtin …]

Cada integración reemplaza los precios del mismo producto y marketplace (el más reciente gana); un precio vacío en el
Excel borra ese renglón. «Precio de venta en otros marketplaces» = el precio más bajo encontrado para el producto y
«Otros marketplaces» = la tienda de ese precio.
"""
import argparse
import csv
import datetime
import glob
import json
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV = os.path.join(ROOT, "data", "precios_otros_marketplaces.csv")
CAMPOS = ["gtin", "marketplace", "precio", "precio_lista", "url", "fecha", "nota", "origen"]
# Nombres canónicos (lo que escriba el dueño o un agente se normaliza a estos)
ALIAS = {
    "benavides": "Farmacias Benavides", "farmacias benavides": "Farmacias Benavides",
    "ahorro": "Farmacias del Ahorro", "del ahorro": "Farmacias del Ahorro", "farmacias del ahorro": "Farmacias del Ahorro", "fahorro": "Farmacias del Ahorro",
    "yza": "Farmacias YZA", "farmacias yza": "Farmacias YZA",
    "san pablo": "Farmacias San Pablo", "farmacias san pablo": "Farmacias San Pablo",
    "guadalajara": "Farmacias Guadalajara", "farmacias guadalajara": "Farmacias Guadalajara",
    "similares": "Farmacias Similares", "farmacias similares": "Farmacias Similares",
    "walmart": "Walmart", "bodega aurrera": "Bodega Aurrera", "aurrera": "Bodega Aurrera",
    "amazon": "Amazon", "amazon mexico": "Amazon", "amazon méxico": "Amazon",
    "mercado libre": "Mercado Libre", "mercadolibre": "Mercado Libre", "meli": "Mercado Libre",
    "chedraui": "Chedraui", "soriana": "Soriana", "heb": "HEB", "h-e-b": "HEB", "la comer": "La Comer",
    "sanborns": "Sanborns", "farmalisto": "Farmalisto", "costco": "Costco", "sams": "Sam's Club", "sam's": "Sam's Club", "sam's club": "Sam's Club",
}


def nombre(m):
    m = (m or "").strip()
    return ALIAS.get(m.lower(), m)


def num(v):
    if v in (None, ""):
        return None
    try:
        return round(float(str(v).replace("$", "").replace(",", "").strip()), 2)
    except ValueError:
        return None


def leer():
    """{gtin: [renglones]} ordenados por precio."""
    out = {}
    if not os.path.exists(CSV):
        return out
    for r in csv.DictReader(open(CSV, encoding="utf-8")):
        r["precio"] = num(r["precio"])
        r["precio_lista"] = num(r.get("precio_lista"))
        if r["precio"] is not None:
            out.setdefault(r["gtin"], []).append(r)
    for v in out.values():
        v.sort(key=lambda r: (r["precio"], r["marketplace"]))
    return out


def resumen(renglones):
    """(precio más bajo, marketplace de ese precio, precio más alto, marketplace, número de marketplaces)."""
    if not renglones:
        return None, "", None, "", 0
    lo, hi = renglones[0], max(renglones, key=lambda r: r["precio"])
    return lo["precio"], lo["marketplace"], hi["precio"], hi["marketplace"], len(renglones)


def escribir(datos):
    orden = [r["gtin"] for r in csv.DictReader(open(os.path.join(ROOT, "data", "prioridad.csv"), encoding="utf-8"))]
    pos = {g: i for i, g in enumerate(orden)}
    filas = [r for v in datos.values() for r in v]
    filas.sort(key=lambda r: (pos.get(r["gtin"], 10 ** 6), r["precio"] if r["precio"] is not None else 0, r["marketplace"]))
    with open(CSV, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=CAMPOS)
        w.writeheader()
        for r in filas:
            w.writerow({k: ("" if r.get(k) is None else (f"{r[k]:.2f}" if k in ("precio", "precio_lista") else r[k])) for k in CAMPOS})


def poner(datos, g, marketplace, precio, precio_lista=None, url="", fecha="", nota="", origen=""):
    m = nombre(marketplace)
    lista = [r for r in datos.get(g, []) if r["marketplace"] != m]
    if precio is not None:
        lista.append({"gtin": g, "marketplace": m, "precio": precio, "precio_lista": precio_lista, "url": url or "", "fecha": fecha,
                      "nota": nota or "", "origen": origen})
    if lista:
        datos[g] = lista
    else:
        datos.pop(g, None)


def integrar(sesion, fecha):
    datos = leer()
    n = 0
    for f in sorted(glob.glob(os.path.join(ROOT, "trabajo", f"sesion_{sesion:02d}", "resultados", "precios_*.json"))):
        for x in json.load(open(f, encoding="utf-8")):
            for p in x.get("precios", []):
                precio = num(p.get("precio"))
                if precio is None or precio <= 0:
                    continue
                lista = num(p.get("precio_lista"))
                poner(datos, str(x["gtin"]), p["marketplace"], precio, lista if lista and lista > precio else None,
                      p.get("url", ""), p.get("fecha") or fecha, p.get("nota", ""), f"agente sesión {sesion}")
                n += 1
    escribir(datos)
    return n, len(datos)


def desde_excel(archivo, fecha):
    import openpyxl
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    ws = wb["Captura"] if "Captura" in wb.sheetnames else wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(c or "").strip().lower() for c in filas[0]]

    def col(*nombres):
        return next((enc.index(n) for n in nombres if n in enc), None)

    cg, cm, cp = col("código", "codigo", "gtin"), col("marketplace", "otros marketplaces"), col("precio", "precio de venta en otros marketplaces")
    cl, cu, cn, cf = col("precio lista", "precio de lista"), col("url", "página"), col("nota"), col("fecha")
    if None in (cg, cm, cp):
        raise SystemExit("El Excel debe tener las columnas Código, Marketplace y Precio (hoja Captura)")
    datos = leer()
    n = 0
    for r in filas[1:]:
        g, m = str(r[cg] or "").strip(), str(r[cm] or "").strip()
        if not g or not m:
            continue
        f = r[cf] if cf is not None else None
        f = f.date().isoformat() if hasattr(f, "date") else (str(f).strip()[:10] if f else fecha)
        poner(datos, g, m, num(r[cp]), num(r[cl]) if cl is not None else None, str(r[cu] or "") if cu is not None else "",
              f, str(r[cn] or "") if cn is not None else "", "dueño (Excel)")
        n += 1
    escribir(datos)
    return n, len(datos)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a1 = sub.add_parser("integrar")
    a1.add_argument("--sesion", type=int, required=True)
    a2 = sub.add_parser("excel")
    a2.add_argument("archivo")
    a3 = sub.add_parser("lista")
    a3.add_argument("--gtin", nargs="*")
    for s in (a1, a2):
        s.add_argument("--fecha", default=datetime.date.today().isoformat())
    a = ap.parse_args()
    if a.cmd == "integrar":
        n, p = integrar(a.sesion, a.fecha)
        print(f"{n} precios integrados; {p} productos con precio en otros marketplaces -> {os.path.relpath(CSV, ROOT)}")
    elif a.cmd == "excel":
        n, p = desde_excel(a.archivo, a.fecha)
        print(f"{n} renglones aplicados; {p} productos con precio en otros marketplaces -> {os.path.relpath(CSV, ROOT)}")
    else:
        for g, v in leer().items():
            if a.gtin and g not in a.gtin:
                continue
            print(g, " · ".join(f"{r['marketplace']} ${r['precio']:,.2f}" for r in v))


if __name__ == "__main__":
    main()
