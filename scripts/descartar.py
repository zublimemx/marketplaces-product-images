#!/usr/bin/env python3
"""Descarta productos del catálogo de importación a Mercado Libre, o los reactiva.

Uso:
  python scripts/descartar.py --gtin 7501... 7502... --motivo "Indicación del dueño: ..."
  python scripts/descartar.py --gtin 7501... --reactivar
  python scripts/descartar.py --lista
  python scripts/descartar.py --excel descartes_AAAA-MM-DD_HHMM.xlsx   # Excel «Exportar descartes» del visor

Marca marketplaces.mercadolibre.descartado = {motivo, fecha} en product.json. El producto se queda en el
repositorio (fichas y fotos sirven para otros marketplaces), pero scripts/build_mercadolibre.py lo saca del
layout y lo lista en la hoja Descartados, y el visor lo muestra como descartado y sin pendientes.
El Excel del visor (hoja Descartes: Código, Producto, Acción = Descartar o Reactivar, Motivo, Fecha) aplica cada
renglón; un descarte sin motivo queda como «Descartado desde el visor».
Después: regenerar el layout y el visor (docs/PROCEDIMIENTO_SESION.md §4).
"""
import argparse
import datetime
import glob
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def ruta(gtin):
    return os.path.join(ROOT, "products", gtin, "product.json")


def leer(gtin):
    f = ruta(gtin)
    if not os.path.exists(f):
        raise SystemExit(f"No existe products/{gtin}/product.json")
    return json.load(open(f, encoding="utf-8"))


def guardar(gtin, data):
    f = ruta(gtin)
    tmp = f + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    os.replace(tmp, f)


def descartar(gtins, motivo, fecha=None):
    fecha = fecha or datetime.date.today().isoformat()
    hechos = []
    for g in gtins:
        p = leer(g)
        p["marketplaces"]["mercadolibre"]["descartado"] = {"motivo": motivo, "fecha": fecha}
        guardar(g, p)
        hechos.append(g)
    return hechos


def reactivar(gtins):
    hechos = []
    for g in gtins:
        p = leer(g)
        if p["marketplaces"]["mercadolibre"].pop("descartado", None) is not None:
            guardar(g, p)
            hechos.append(g)
    return hechos


def desde_excel(archivo):
    import openpyxl
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    ws = wb["Descartes"] if "Descartes" in wb.sheetnames else wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(c or "").strip().lower() for c in filas[0]]
    col = {k: enc.index(k) for k in ("código", "acción", "motivo", "fecha") if k in enc}
    if "código" not in col or "acción" not in col:
        raise SystemExit("El Excel debe tener las columnas Código y Acción (hoja Descartes)")
    desc, react = [], []
    for r in filas[1:]:
        g = str(r[col["código"]] or "").strip()
        acc = str(r[col["acción"]] or "").strip().lower()
        if not g:
            continue
        if acc.startswith("reactivar"):
            react.append(g)
        elif acc.startswith("descartar"):
            motivo = str(r[col["motivo"]] or "").strip() if "motivo" in col else ""
            fecha = r[col["fecha"]] if "fecha" in col else None
            fecha = fecha.date().isoformat() if hasattr(fecha, "date") else (str(fecha).strip()[:10] if fecha else None)
            desc.append((g, motivo or "Descartado desde el visor", fecha))
    for g, motivo, fecha in desc:
        descartar([g], motivo, fecha)
    return [g for g, _, _ in desc], reactivar(react)


def listar():
    out = []
    for f in sorted(glob.glob(os.path.join(ROOT, "products", "*", "product.json"))):
        p = json.load(open(f, encoding="utf-8"))
        d = p["marketplaces"]["mercadolibre"].get("descartado")
        if d:
            out.append((p["gtin"], p.get("titulo") or p["nombre_sistema"], d["fecha"], d["motivo"]))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--gtin", nargs="+", default=[])
    ap.add_argument("--motivo", help="obligatorio al descartar; di quién lo pidió y por qué")
    ap.add_argument("--fecha", help="AAAA-MM-DD; por omisión, hoy")
    ap.add_argument("--reactivar", action="store_true", help="quita el descarte y el producto vuelve al layout")
    ap.add_argument("--lista", action="store_true", help="lista los productos descartados")
    ap.add_argument("--excel", help="Excel «Exportar descartes» del visor")
    a = ap.parse_args()
    if a.excel:
        d, r = desde_excel(a.excel)
        print(f"Descartados: {len(d)} · reactivados: {len(r)}")
        print("Siguiente: regenera el layout y el visor (docs/PROCEDIMIENTO_SESION.md §4).", file=sys.stderr)
        return
    if a.lista:
        filas = listar()
        for g, t, f, m in filas:
            print(f"{g}\t{f}\t{t}\t{m}")
        print(f"{len(filas)} descartados")
        return
    if not a.gtin:
        ap.error("indica --gtin o --excel, o usa --lista")
    if a.reactivar:
        hechos = reactivar(a.gtin)
        print(f"Reactivados: {len(hechos)} de {len(a.gtin)}")
    else:
        if not a.motivo:
            ap.error("--motivo es obligatorio al descartar")
        hechos = descartar(a.gtin, a.motivo, a.fecha)
        print(f"Descartados: {len(hechos)}")
    print("Siguiente: regenera el layout y el visor (docs/PROCEDIMIENTO_SESION.md §4).", file=sys.stderr)


if __name__ == "__main__":
    main()
